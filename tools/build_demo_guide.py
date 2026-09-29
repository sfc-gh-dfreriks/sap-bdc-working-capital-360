#!/usr/bin/env python3
"""Build docs/SAP_Working_Capital_360_Demo_Guide.pptx.

Uses the sibling Spend 360 demo guide as the template (same Snowflake layouts and
shape positions). Each text frame is rewritten in place, keeping its run formatting,
and three screenshot slides are added after Demo Act 1. Figures come from
/tmp/wc_facts.json (run tools/wc_facts.py); screenshots come from
/tmp/wc_shots_kit (run tools/capture_shots.py).
    python3 tools/build_demo_guide.py
"""
import copy
import json
import pathlib

from pptx import Presentation
from pptx.util import Emu

ROOT = pathlib.Path(__file__).resolve().parents[1]
TEMPLATE = ROOT.parent / "sap-bdc-spend-360/docs/SAP_Spend_360_Demo_Guide.pptx"
OUT = ROOT / "docs/SAP_Working_Capital_360_Demo_Guide.pptx"
SHOTS = pathlib.Path("/tmp/wc_shots_kit")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
F = json.load(open("/tmp/wc_facts.json"))["facts"]

L, P = F["kpi_total_latest"], F["kpi_total_prior"]
HI = F["highest_ccc_company"]
CASH_M = F["cash_release_total_usd"] / 1e6
HONEST = ("AR/AP invoices, amounts and dates are real BDC Entry View Journal Entry lines; "
          "payment terms and behaviour, early-pay programs, inventory, bank balances and partner "
          "names are demo enrichment (IS_DEMO_ENRICHMENT). USD at fixed illustrative FX (EUR 1.08, "
          "JPY 0.0067). NWC is negative because AP is ~3x AR in the demo tenant.")


def set_text(shape, *texts):
    """Replace the non-empty paragraphs of a shape, in order, keeping formatting."""
    paras = [p for p in shape.text_frame.paragraphs if p.text.strip()]
    for p, t in zip(paras, texts):
        runs = p.runs
        runs[0].text = t
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        for br in p._p.findall(f"{{{A}}}br"):
            p._p.remove(br)
        if "\v" in t:  # soft line break: run, <a:br/>, cloned run
            first, rest = t.split("\v", 1)
            runs[0].text = first
            br = copy.deepcopy(runs[0]._r); br.tag = f"{{{A}}}br"
            for c in list(br):
                if c.tag != f"{{{A}}}rPr":
                    br.remove(c)
            r2 = copy.deepcopy(runs[0]._r)
            runs[0]._r.addnext(br); br.addnext(r2)
            r2.find(f"{{{A}}}t").text = rest
    # drop surplus non-empty paragraphs
    for p in paras[len(texts):]:
        p._p.getparent().remove(p._p)


# (slide index 1-based) -> {shape index: texts}, notes
CONTENT = {
    1: ({0: ["Demo Guide & the Snowflake × SAP Business Data Cloud story"],
         1: ["Dave Freriks  |  Sales Engineering"],
         2: ["SAP WORKING CAPITAL 360\vON SNOWFLAKE"]},
        "Welcome. Today I'm showing SAP Working Capital 360, a Snowflake Native App built on SAP "
        "Business Data Cloud finance data. It mirrors SAP Working Capital Insights and the SAP Taulia "
        "early-payment levers. The story: the CFO asks why our cash conversion cycle grew year over year."),
    2: ({0: ["Why Snowflake + SAP Business Data Cloud"],
         1: ["Receivables, payables, cash and inventory live in SAP, but a working capital view needs "
             "them together with bank balances and funding programs. That usually means weeks of ETL, "
             "and the business context is lost along the way."],
         2: ["SAP BDC Connect shares governed journal entry, invoice and cash data products into "
             "Snowflake with zero copy. Snowflake joins non-SAP data and adds AI. No pipelines are "
             "needed and the context is preserved."],
         3: ["The Challenge"], 4: ["The Opportunity"],
         6: ["Two platforms, one governed data foundation"]},
        "Set up the problem. DSO, DPO and DIO sit in different SAP modules and the cash position sits "
        "at the banks. BDC shares the SAP side zero-copy; Snowflake joins it with non-SAP data and adds AI."),
    3: ({0: ["What Is Working Capital 360?"],
         1: ["9 Working Capital Insights pages: cash conversion cycle, cash, AR, AP, early payment, inventory, levers."],
         2: ["A Cortex Agent answers plain-English questions over a governed semantic view."],
         3: ["Runs on SPCS with data, semantic view and UI bundled; nothing leaves Snowflake."],
         5: ["A self-contained Snowflake Native App built on SAP BDC data"],
         9: ["Dashboards"], 10: ["Ask the Agent"], 11: ["Runs In Snowflake"]},
        "Tell them what they're about to see: 9 pages spelling out DSO (Days Sales Outstanding), DPO "
        "(Days Payables Outstanding), DIO (Days Inventory Outstanding) and CCC (cash conversion cycle = DSO + DIO − DPO), plus a live Cortex agent."),
    4: ({0: ["How It’s Built: BDC → Snowflake → App"],
         1: ["SAP BDC Connect shares 11 finance data products, including Entry View Journal Entry, "
             "Supplier Invoice and Cash Flow, with zero copy and no ETL."],
         2: ["14 L1 views feed L2 dynamic tables with DSO, DPO, DIO and CCC per company and month, "
             "plus AR/AP items, inventory and a 13-week cash forecast."],
         3: ["A Native App packages the data, the semantic view and an SPCS React UI with a Cortex "
             "agent. It installs in one click."],
         4: ["1. BDC Data Products"], 5: ["2. Snowflake Curation"], 6: ["3. Native App + Cortex"],
         8: ["From SAP data product to a shareable AI app"]},
        "Walk the three steps: zero-copy in, curated into Working Capital Insights KPIs with the Taulia "
        "levers, and packaged as a one-click app. The dynamic tables refresh daily on their own."),
    5: ({0: ["DEMOING WORKING CAPITAL 360", "THE 10-MINUTE FLOW"]},
        "Transition: 'Let me show you the app. The CFO wants to know why our cash conversion cycle "
        "grew year over year.'"),
    6: ({0: ["The Demo Flow"],
         1: ["Open the app: no setup, the data is already inside",
             f"Overview: cash cycle (CCC) {P['CCC']} → {L['CCC']} days",
             f"AP and Early Payment: DPO fell {P['DPO']} → {L['DPO']} days",
             "AR and Inventory: overdue, disputes, DIO vs target",
             f"WC Opportunities: USD {CASH_M:.1f}M cash release by lever",
             "Ask the Agent why Japan's cash cycle rose",
             "Recap: zero-ETL, governed, AI-ready"]},
        "Roadmap for the live demo; don't read it aloud. Focus most on AP and Early Payment & "
        "SCF: shorter DPO is the main reason CCC grew, and the Taulia levers fix it."),
    7: ({0: ["Demo Act 1: The Dashboards"],
         1: [f"Cash conversion cycle (DSO + DIO − DPO) rose {F['ccc_change_days']} days to {L['CCC']}; "
             f"{HI['COMPANY']} is highest at {HI['CCC']} days."],
         2: [f"DPO fell from {P['DPO']} to {L['DPO']} days; USD {L['DISC_LOST_USD']/1e3:.1f}K of "
             "discounts were lost last month."],
         3: [f"Dynamic discounting (APR {F['dd_effective_apr']['APR_PCT']}%) and supply chain finance "
             "are the Taulia levers."],
         4: [f"USD {CASH_M:.1f}M cash release identified; SCF-extended DPO is the largest lever."],
         6: ["Answer the CFO: why did the cash conversion cycle grow?"],
         10: ["Overview"], 11: ["Payables"], 12: ["Early Payment"], 14: ["Opportunities"]},
        "Click through and narrate outcomes: 'CCC is up 4.7 days, and the driver is DPO: we pay "
        "suppliers earlier and still lose discounts.' Then show the SCF and dynamic discounting levers."),
    8: ({0: ["Demo Act 2: Ask the Agent"],
         1: ["Ask in plain English: “Why did the cash conversion cycle increase in Japan in 2025?”, “How much "
             "discount did we lose by program?” or “Which customers have the most overdue AR?”"],
         2: ["SAP_WORKING_CAPITAL_ANALYST maps the question to the bundled semantic view, generates "
             "governed SQL and returns a chart with the exact query, so the answer is explainable."],
         4: ["Cortex Agent over a governed semantic view"]},
        "The highlight. Type a question live; do NOT pre-load it. Expand the generated SQL to show "
        "the answer is governed and explainable."),
    9: ({0: ["Lead with the CFO question, not architecture. AR/AP invoices are real BDC journal lines; "
             "terms, early-pay programs, inventory, bank balances and partner names are demo "
             "enrichment. USD at fixed FX (EUR 1.08, JPY 0.0067); NWC is negative as AP is ~3x AR."],
         1: ["Key Points to Land During the Demo"], 3: ["Say these out loud"]},
        "Reinforce: no ETL, Working Capital Insights plus Taulia levers, AI-ready and governed. Always "
        "state the data honesty line; it builds trust with finance audiences."),
    10: ({0: ["The Snowflake × SAP BDC Value"],
          1: ["BDC Connect zero-copy sharing means no pipelines; journal entry data is live in Snowflake."],
          2: ["SAP semantics are preserved and joined with non-SAP bank and funding data."],
          3: ["A Cortex Agent brings natural-language analytics directly to SAP working capital data."],
          4: ["Native Apps package data + AI + UI into a one-click, distributable product."],
          6: ["Why this partnership matters"],
          10: ["Zero-ETL"], 11: ["Governed"], 12: ["AI-Ready"], 14: ["Global Apps"]},
         "Zoom out to the partnership value. Cover each of the four pillars in one sentence."),
    11: ({0: ["Takeaways by the Numbers"],
          1: ["ETL pipelines to build: BDC Connect shares data zero-copy."],
          2: ["BDC finance data products composed into one app."],
          3: ["Days of cash conversion cycle growth, explained to the driver."],
          4: ["Cash release identified across the SAP Taulia levers."],
          5: ["0"], 6: ["11"], 7: [str(F["ccc_change_days"])], 8: [f"${CASH_M:.0f}M"],
          10: ["What this demo proves"]},
         f"Make it concrete: zero pipelines, 11 BDC data products, a {F['ccc_change_days']}-day CCC "
         f"increase explained, and USD {CASH_M:.1f}M of cash release identified."),
    12: ({21: ["One App. US Live. EMEA & APAC Next."]},
         "The app is live in US (AWS_US_WEST_2) as ORGDATACLOUD$INTERNAL$WORKING_CAPITAL_360_ORG. "
         "EMEA and APAC use the same scripts with --target dfreriks_eu_demo / dfreriks_apac_demo; "
         "each region is its own governed install."),
    13: ({0: ["Demo Do’s and Don’ts"],
          1: ["Lead with the CFO question. Show the app already has data. Focus on AP and Early "
              "Payment & SCF. Ask SAP_WORKING_CAPITAL_ANALYST a real question live."],
          2: ["Don’t pre-load canned answers or dwell on SQL. Don’t present enriched terms, programs, "
              "inventory or bank balances as real tenant data. Don't hide the negative NWC; explain it."],
          3: ["Do"], 4: ["Don’t"], 6: ["Keep it crisp and outcome-focused"]},
         "Presenter coaching. " + HONEST),
    14: ({0: ["Next Steps"],
          1: ["For SEs: install from the internal marketplace (or use the public demo) and rehearse "
              "the 10-minute flow."],
          2: ["For customers: map SAP finance data products in BDC and pick the first working "
              "capital lever to size."],
          3: ["For partners: co-build Native Apps on BDC data with SAP Taulia levers to reach every "
              "Snowflake region."],
          5: ["Take the story forward"]},
         "Close with a clear ask per audience. Public demo: "
         "https://sfc-gh-dfreriks.github.io/working-capital-360-public/"),
    15: ({}, "Thank the audience. Share the listing locator ORGDATACLOUD$INTERNAL$WORKING_CAPITAL_360_ORG "
             "(role WORKING_CAPITAL_360_APP_USERS) so they can install Working Capital 360."),
}

SHOT_SLIDES = [
    ("The Dashboards: CCC, Payables & Early Payment",
     ["overview", "ap", "early-pay", "ar"],
     "Screens for Act 1: Working Capital Overview, Accounts Payable, Early Payment & SCF, Accounts Receivable."),
    ("The Dashboards: Inventory, Cash & Opportunities",
     ["inventory", "cash", "opportunities", "lineage"],
     "Inventory DIO vs target, Cash & Liquidity with the 13-week forecast, WC Opportunities, and BDC Sources & Lineage."),
    ("Ask the Agent", ["analyst"],
     "The Ask the Agent page. Ask a live question and expand the SQL."),
]


def add_shot_slide(prs, title_src, title, shots, notes):
    s = prs.slides.add_slide(title_src.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    # copy the title box from the template slide so fonts match
    tbox = copy.deepcopy(title_src.shapes[0]._element)
    s.shapes._spTree.append(tbox)
    set_text(s.shapes[-1], title)
    W, H = prs.slide_width, prs.slide_height
    top, margin, gap = Emu(int(H * 0.2)), Emu(int(W * 0.05)), Emu(int(W * 0.015))
    if len(shots) == 1:
        h = H - top - Emu(int(H * 0.05)); w = int(h * 1.6)
        s.shapes.add_picture(str(SHOTS / f"{shots[0]}.png"), (W - w) // 2, top, w, h)
    else:
        avail_h = H - top - Emu(int(H * 0.04))
        h = (avail_h - gap) // 2; w = int(h * 1.6)
        x0 = (W - 2 * w - gap) // 2
        for i, sid in enumerate(shots):
            s.shapes.add_picture(str(SHOTS / f"{sid}.png"), x0 + (i % 2) * (w + gap),
                                 top + (i // 2) * (h + gap), w, h)
    s.notes_slide.notes_text_frame.text = notes
    return s


def main():
    prs = Presentation(TEMPLATE)
    slides = list(prs.slides)
    for idx, (shapes, notes) in CONTENT.items():
        sl = slides[idx - 1]
        for j, texts in shapes.items():
            set_text(sl.shapes[j], *texts)
        sl.notes_slide.notes_text_frame.text = notes
    # screenshot slides after Act 1 (slide 7), styled like slide 8 (One Column Layout_1)
    ids = prs.slides._sldIdLst
    new = [add_shot_slide(prs, slides[7], *spec) for spec in SHOT_SLIDES]
    elems = list(ids)[-len(new):]
    for e in elems:
        ids.remove(e)
    for k, e in enumerate(elems):
        ids.insert(7 + k, e)
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
