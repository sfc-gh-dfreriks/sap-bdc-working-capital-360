"""Layout helpers shared by the Working Capital 360 decks. Copied from the People 360 deck
builder (lines 55-144) so the kits lay out identically; see that file for history."""
import json
import pathlib
import sys

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from pptx_kit import (  # noqa: E402
    BODY_GREY, DK1, DK2, LIGHT_BG, SF_BLUE, TEAL, VIOLET, WHITE,
    add_shape_text, new_presentation, set_ph, verify_deck, verify_slide,
)
from pptx.dml.color import RGBColor  # noqa: E402
RED = RGBColor(0xA2, 0x00, 0x00)
TOP, BOTTOM, LEFT, RIGHT = 1.32, 5.08, 0.40, 9.50
FULLW = RIGHT - LEFT



def content(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    set_ph(s, 0, title)
    set_ph(s, 1, subtitle)
    return s


def box(slide, x, y, w, h, fill=LIGHT_BG, accent=None):
    add_shape_text(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, "", fill, DK1)
    if accent is not None:
        add_shape_text(slide, MSO_SHAPE.RECTANGLE, x, y, 0.05, h, "", accent, DK1)


def stack(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, spacing=1.06):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(runs):
        body, size, bold, colour = item[:4]
        after = item[4] if len(item) > 4 else 5
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = body
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Arial"
    return tb


def card(slide, x, y, w, h, kicker, runs, accent=DK2, fill=LIGHT_BG):
    box(slide, x, y, w, h, fill, accent)
    head = [(kicker.upper(), 8.5, True, DK2, 7)] if kicker else []
    stack(slide, x + 0.24, y + 0.16, w - 0.42, h - 0.30, head + list(runs))


def banner(slide, y, runs, fill=DK2, h=0.62, x=LEFT, w=FULLW):
    add_shape_text(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, "", fill, WHITE)
    stack(slide, x + 0.30, y + 0.11, w - 0.60, h - 0.22, runs, align=PP_ALIGN.LEFT)


def stat(slide, x, y, w, h, value, label, detail=None, fill=LIGHT_BG, accent=DK2):
    box(slide, x, y, w, h, fill, accent)
    runs = [(value, 28, True, DK2, 3), (label.upper(), 8.5, True, BODY_GREY, 4)]
    if detail:
        runs.append((detail, 10, False, DK1, 0))
    stack(slide, x + 0.26, y + 0.16, w - 0.46, h - 0.30, runs)


def note(slide, text, y=4.88):
    stack(slide, LEFT, y, FULLW, 0.19, [(text, 8, False, BODY_GREY, 0)])


def picture(slide, path, x, y, w, h):
    im = Image.open(path)
    ar = im.width / im.height
    if ar > w / h:
        pw, ph = w, w / ar
    else:
        ph, pw = h, h * ar
    slide.shapes.add_picture(
        str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
        Inches(pw), Inches(ph))


def caption(slide, x, y, w, text):
    stack(slide, x, y, w, 0.20, [(text.upper(), 8.5, True, DK2, 0)])


def region_label(raw: str) -> str:
    """AWS us-west-2 rather than PUBLIC.AWS_US_WEST_2.

    The raw identifier is a long all-caps token: the deck verifier flags it as
    merged caps, and it reads badly at 12pt. Cloud stays capitalised, the region
    goes lower-case with hyphens, which is how the provider writes it anyway.
    """
    name = raw.split(".")[-1]
    parts = name.split("_")
    if not parts:
        return raw
    cloud, rest = parts[0], parts[1:]
    return f"{cloud} {'-'.join(x.lower() for x in rest)}" if rest else cloud

